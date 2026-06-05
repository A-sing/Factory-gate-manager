"""DBS FACTORY - Gate Management System Backend.

FastAPI + MongoDB + JWT auth (admin / guard roles).
All routes are mounted under /api so they flow through the ingress.
"""
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Optional
import logging
import os
import re
import uuid

import jwt
from bson import ObjectId  # noqa: F401 (kept for future use)
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, HTTPException, Query, status
from fastapi.responses import Response
from fastapi.security import OAuth2PasswordBearer
from motor.motor_asyncio import AsyncIOMotorClient
from passlib.context import CryptContext
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware


# ---------------------------------------------------------------------------
# Config / globals
# ---------------------------------------------------------------------------

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ.get("DB_NAME", "dbs_factory")
JWT_SECRET = os.environ.get("JWT_SECRET", "dbs-factory-dev-secret-change-me")
JWT_ALGORITHM = "HS256"
JWT_EXPIRE_HOURS = 12

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login", auto_error=True)

app = FastAPI(title="DBS Factory Gate Management")
api = FastAPI()  # not used; routes below use APIRouter via app

from fastapi import APIRouter  # noqa: E402

api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("dbs_factory")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def hash_password(p: str) -> str:
    return pwd_context.hash(p)


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return pwd_context.verify(plain, hashed)
    except Exception:
        return False


def create_token(user: dict) -> str:
    now = utcnow()
    payload = {
        "sub": user["id"],
        "username": user["username"],
        "role": user["role"],
        "name": user.get("name", user["username"]),
        "iat": int(now.timestamp()),
        "exp": int((now + timedelta(hours=JWT_EXPIRE_HOURS)).timestamp()),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(token: str = Depends(oauth2_scheme)) -> dict:
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid token")
    user = await db.users.find_one({"id": payload.get("sub")}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "User not found")
    return user


def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Admin only")
    return user


async def ensure_indexes() -> None:
    await db.users.create_index("username", unique=True)
    await db.labours.create_index("labour_id", unique=True)
    await db.counters.create_index("name", unique=True)


async def next_labour_id() -> str:
    res = await db.counters.find_one_and_update(
        {"name": "labour"},
        {"$inc": {"value": 1}},
        upsert=True,
        return_document=True,
    )
    n = res["value"] if res else 1
    return f"LAB-{n:06d}"


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------

class LoginRequest(BaseModel):
    username: str
    password: str


class ChangePasswordRequest(BaseModel):
    current_password: str
    new_password: str


class UserCreate(BaseModel):
    name: str
    username: str
    password: str
    role: str  # "admin" or "guard"


class UserUpdate(BaseModel):
    name: Optional[str] = None
    password: Optional[str] = None
    role: Optional[str] = None


class SettingsUpdate(BaseModel):
    business_name: Optional[str] = None
    labour_categories: Optional[List[str]] = None
    visit_purposes: Optional[List[str]] = None
    gates: Optional[List[str]] = None


class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: dict


class UserPublic(BaseModel):
    id: str
    name: str
    username: str
    role: str


class VisitorIn(BaseModel):
    visitor_name: str
    mobile_number: Optional[str] = None
    purpose: str
    photo_base64: Optional[str] = None
    gate_name: Optional[str] = None


class Visitor(BaseModel):
    id: str
    visitor_name: str
    mobile_number: Optional[str] = None
    purpose: str
    photo_base64: Optional[str] = None
    gate_name: Optional[str] = None
    entry_datetime: datetime
    created_by: str
    created_by_name: str


class ContractorIn(BaseModel):
    contractor_name: str
    contact_person: Optional[str] = None
    mobile_number: Optional[str] = None
    address: Optional[str] = None


class Contractor(ContractorIn):
    id: str
    created_at: datetime


class LabourIn(BaseModel):
    labour_name: str
    contractor_id: str
    category: str
    photo_base64: str
    aadhaar_number: Optional[str] = None
    mobile_number: Optional[str] = None


class Labour(BaseModel):
    id: str
    labour_id: str
    labour_name: str
    contractor_id: str
    contractor_name: Optional[str] = None
    category: str
    photo_base64: str
    aadhaar_number: Optional[str] = None
    mobile_number: Optional[str] = None
    created_at: datetime


class AttendanceRecord(BaseModel):
    id: str
    labour_id: str
    labour_name: str
    contractor_name: Optional[str] = None
    category: str
    check_in_time: datetime
    check_out_time: Optional[datetime] = None
    total_hours: Optional[float] = None
    gate_name: Optional[str] = None
    status: str  # "inside" or "checked_out"


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

@api_router.post("/auth/login", response_model=TokenResponse)
async def login(body: LoginRequest):
    user = await db.users.find_one({"username": body.username.lower().strip()})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid username or password")
    public = {
        "id": user["id"],
        "name": user["name"],
        "username": user["username"],
        "role": user["role"],
    }
    token = create_token(public)
    return TokenResponse(access_token=token, user=public)


@api_router.get("/auth/me", response_model=UserPublic)
async def me(user: dict = Depends(get_current_user)):
    return UserPublic(**{k: user[k] for k in ("id", "name", "username", "role")})


@api_router.post("/auth/change-password")
async def change_password(body: ChangePasswordRequest, user: dict = Depends(get_current_user)):
    full = await db.users.find_one({"id": user["id"]})
    if not full or not verify_password(body.current_password, full["password_hash"]):
        raise HTTPException(400, "Current password is incorrect")
    if len(body.new_password) < 6:
        raise HTTPException(400, "New password must be at least 6 characters")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": hash_password(body.new_password)}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# User management (admin only)
# ---------------------------------------------------------------------------

@api_router.get("/users", response_model=List[UserPublic])
async def list_users(_: dict = Depends(require_admin)):
    cursor = db.users.find({}, {"_id": 0, "password_hash": 0}).sort("created_at", 1)
    return [UserPublic(**{k: d[k] for k in ("id", "name", "username", "role")}) async for d in cursor]


@api_router.post("/users", response_model=UserPublic)
async def create_user(body: UserCreate, _: dict = Depends(require_admin)):
    uname = body.username.lower().strip()
    if not uname or not body.password or not body.name.strip():
        raise HTTPException(400, "Name, username and password are required")
    if body.role not in ("admin", "guard"):
        raise HTTPException(400, "Role must be admin or guard")
    if len(body.password) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    if await db.users.find_one({"username": uname}):
        raise HTTPException(400, "Username already exists")
    doc = {
        "id": str(uuid.uuid4()),
        "name": body.name.strip(),
        "username": uname,
        "password_hash": hash_password(body.password),
        "role": body.role,
        "created_at": utcnow(),
    }
    await db.users.insert_one(doc)
    return UserPublic(id=doc["id"], name=doc["name"], username=doc["username"], role=doc["role"])


@api_router.put("/users/{uid}", response_model=UserPublic)
async def update_user(uid: str, body: UserUpdate, _: dict = Depends(require_admin)):
    update: dict = {}
    if body.name is not None:
        update["name"] = body.name.strip()
    if body.role is not None:
        if body.role not in ("admin", "guard"):
            raise HTTPException(400, "Role must be admin or guard")
        update["role"] = body.role
    if body.password is not None:
        if len(body.password) < 6:
            raise HTTPException(400, "Password must be at least 6 characters")
        update["password_hash"] = hash_password(body.password)
    if not update:
        raise HTTPException(400, "No changes")
    res = await db.users.find_one_and_update(
        {"id": uid}, {"$set": update}, return_document=True, projection={"_id": 0, "password_hash": 0}
    )
    if not res:
        raise HTTPException(404, "User not found")
    return UserPublic(**{k: res[k] for k in ("id", "name", "username", "role")})


@api_router.delete("/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_admin)):
    if uid == user["id"]:
        raise HTTPException(400, "You cannot delete yourself")
    res = await db.users.delete_one({"id": uid})
    if res.deleted_count == 0:
        raise HTTPException(404, "User not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Settings (business name + dropdowns)
# ---------------------------------------------------------------------------

DEFAULT_SETTINGS = {
    "business_name": "DBS Factory",
    "labour_categories": ["Welder", "Fitter", "Electrician", "Helper", "Carpenter", "Painter", "Other"],
    "visit_purposes": ["Meeting", "Delivery", "Maintenance", "Interview", "Audit", "Other"],
    "gates": ["Main Gate", "Gate 1", "Gate 2"],
}


async def _settings_doc() -> dict:
    doc = await db.settings.find_one({"key": "app"}, {"_id": 0})
    if not doc:
        doc = {"key": "app", **DEFAULT_SETTINGS, "updated_at": utcnow()}
        await db.settings.insert_one(doc)
        doc.pop("_id", None)
    for k, v in DEFAULT_SETTINGS.items():
        doc.setdefault(k, v)
    return doc


def _public_settings(doc: dict) -> dict:
    return {
        "business_name": doc["business_name"],
        "labour_categories": doc["labour_categories"],
        "visit_purposes": doc["visit_purposes"],
        "gates": doc["gates"],
    }


@api_router.get("/settings")
async def get_settings(_: dict = Depends(get_current_user)):
    doc = await _settings_doc()
    return _public_settings(doc)


@api_router.put("/settings")
async def put_settings(body: SettingsUpdate, _: dict = Depends(require_admin)):
    await _settings_doc()
    update: dict = {"updated_at": utcnow()}
    if body.business_name is not None and body.business_name.strip():
        update["business_name"] = body.business_name.strip()
    if body.labour_categories is not None:
        update["labour_categories"] = [c.strip() for c in body.labour_categories if c.strip()]
    if body.visit_purposes is not None:
        update["visit_purposes"] = [c.strip() for c in body.visit_purposes if c.strip()]
    if body.gates is not None:
        update["gates"] = [c.strip() for c in body.gates if c.strip()]
    await db.settings.update_one({"key": "app"}, {"$set": update})
    return _public_settings(await _settings_doc())




# ---------------------------------------------------------------------------
# Visitors
# ---------------------------------------------------------------------------

@api_router.post("/visitors", response_model=Visitor)
async def create_visitor(body: VisitorIn, user: dict = Depends(get_current_user)):
    doc = {
        "id": str(uuid.uuid4()),
        **body.dict(),
        "entry_datetime": utcnow(),
        "created_by": user["id"],
        "created_by_name": user["name"],
    }
    await db.visitors.insert_one(doc)
    doc.pop("_id", None)
    return Visitor(**doc)


@api_router.get("/visitors", response_model=List[Visitor])
async def list_visitors(
    start: Optional[str] = None,
    end: Optional[str] = None,
    visitor_name: Optional[str] = None,
    mobile_number: Optional[str] = None,
    purpose: Optional[str] = None,
    limit: int = Query(500, le=2000),
    _: dict = Depends(get_current_user),
):
    q: dict = {}
    s_dt = _parse_dt(start)
    e_dt = _parse_dt(end)
    if s_dt or e_dt:
        q["entry_datetime"] = {}
        if s_dt:
            q["entry_datetime"]["$gte"] = s_dt
        if e_dt:
            q["entry_datetime"]["$lte"] = e_dt
    if visitor_name:
        q["visitor_name"] = {"$regex": re.escape(visitor_name), "$options": "i"}
    if mobile_number:
        q["mobile_number"] = {"$regex": re.escape(mobile_number)}
    if purpose:
        q["purpose"] = {"$regex": re.escape(purpose), "$options": "i"}
    cursor = db.visitors.find(q, {"_id": 0}).sort("entry_datetime", -1).limit(limit)
    return [Visitor(**d) async for d in cursor]


# ---------------------------------------------------------------------------
# Contractors
# ---------------------------------------------------------------------------

@api_router.post("/contractors", response_model=Contractor)
async def create_contractor(body: ContractorIn, _: dict = Depends(require_admin)):
    doc = {"id": str(uuid.uuid4()), **body.dict(), "created_at": utcnow()}
    await db.contractors.insert_one(doc)
    doc.pop("_id", None)
    return Contractor(**doc)


@api_router.get("/contractors", response_model=List[Contractor])
async def list_contractors(_: dict = Depends(get_current_user)):
    cursor = db.contractors.find({}, {"_id": 0}).sort("contractor_name", 1)
    return [Contractor(**d) async for d in cursor]


@api_router.put("/contractors/{cid}", response_model=Contractor)
async def update_contractor(cid: str, body: ContractorIn, _: dict = Depends(require_admin)):
    res = await db.contractors.find_one_and_update(
        {"id": cid},
        {"$set": body.dict()},
        return_document=True,
        projection={"_id": 0},
    )
    if not res:
        raise HTTPException(404, "Contractor not found")
    return Contractor(**res)


@api_router.delete("/contractors/{cid}")
async def delete_contractor(cid: str, _: dict = Depends(require_admin)):
    res = await db.contractors.delete_one({"id": cid})
    if res.deleted_count == 0:
        raise HTTPException(404, "Contractor not found")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Labours
# ---------------------------------------------------------------------------

async def _attach_contractor_name(doc: dict) -> dict:
    c = await db.contractors.find_one({"id": doc.get("contractor_id")}, {"_id": 0, "contractor_name": 1})
    doc["contractor_name"] = c["contractor_name"] if c else None
    return doc


@api_router.post("/labours", response_model=Labour)
async def create_labour(body: LabourIn, _: dict = Depends(get_current_user)):
    # uniqueness: if aadhaar provided, ensure not duplicated
    if body.aadhaar_number:
        existing = await db.labours.find_one({"aadhaar_number": body.aadhaar_number})
        if existing:
            raise HTTPException(400, f"Labour with this Aadhaar already exists ({existing['labour_id']})")
    labour_id = await next_labour_id()
    doc = {
        "id": str(uuid.uuid4()),
        "labour_id": labour_id,
        **body.dict(),
        "created_at": utcnow(),
    }
    await db.labours.insert_one(doc)
    doc.pop("_id", None)
    doc = await _attach_contractor_name(doc)
    return Labour(**doc)


@api_router.get("/labours", response_model=List[Labour])
async def list_labours(
    q: Optional[str] = None,
    contractor_id: Optional[str] = None,
    category: Optional[str] = None,
    limit: int = Query(500, le=2000),
    _: dict = Depends(get_current_user),
):
    flt: dict = {}
    if contractor_id:
        flt["contractor_id"] = contractor_id
    if category:
        flt["category"] = category
    if q:
        rx = {"$regex": re.escape(q), "$options": "i"}
        flt["$or"] = [
            {"labour_id": rx},
            {"labour_name": rx},
            {"aadhaar_number": rx},
            {"mobile_number": rx},
        ]
    cursor = db.labours.find(flt, {"_id": 0}).sort("created_at", -1).limit(limit)
    out = []
    async for d in cursor:
        out.append(Labour(**(await _attach_contractor_name(d))))
    return out


@api_router.get("/labours/{labour_id}", response_model=Labour)
async def get_labour(labour_id: str, _: dict = Depends(get_current_user)):
    d = await db.labours.find_one({"labour_id": labour_id}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Labour not found")
    d = await _attach_contractor_name(d)
    return Labour(**d)


@api_router.post("/labours/{labour_id}/checkin", response_model=AttendanceRecord)
async def check_in(labour_id: str, gate_name: Optional[str] = None, user: dict = Depends(get_current_user)):
    labour = await db.labours.find_one({"labour_id": labour_id}, {"_id": 0})
    if not labour:
        raise HTTPException(404, "Labour not found")
    open_rec = await db.attendance.find_one({"labour_id": labour_id, "check_out_time": None})
    if open_rec:
        raise HTTPException(400, "Labour is already inside the factory")
    contractor = await db.contractors.find_one({"id": labour["contractor_id"]}, {"_id": 0})
    rec = {
        "id": str(uuid.uuid4()),
        "labour_id": labour_id,
        "labour_name": labour["labour_name"],
        "contractor_name": contractor["contractor_name"] if contractor else None,
        "category": labour["category"],
        "check_in_time": utcnow(),
        "check_out_time": None,
        "total_hours": None,
        "gate_name": gate_name,
        "status": "inside",
        "guard_id": user["id"],
        "guard_name": user["name"],
    }
    await db.attendance.insert_one(rec)
    rec.pop("_id", None)
    return AttendanceRecord(**rec)


@api_router.post("/labours/{labour_id}/checkout", response_model=AttendanceRecord)
async def check_out(labour_id: str, gate_name: Optional[str] = None, user: dict = Depends(get_current_user)):
    rec = await db.attendance.find_one({"labour_id": labour_id, "check_out_time": None})
    if not rec:
        raise HTTPException(400, "Labour is not currently inside")
    now = utcnow()
    check_in_time = rec["check_in_time"]
    if check_in_time.tzinfo is None:
        check_in_time = check_in_time.replace(tzinfo=timezone.utc)
    hours = round((now - check_in_time).total_seconds() / 3600, 2)
    await db.attendance.update_one(
        {"id": rec["id"]},
        {"$set": {
            "check_out_time": now,
            "total_hours": hours,
            "status": "checked_out",
            "checkout_gate": gate_name,
            "checkout_guard_id": user["id"],
            "checkout_guard_name": user["name"],
        }},
    )
    rec = await db.attendance.find_one({"id": rec["id"]}, {"_id": 0})
    return AttendanceRecord(**rec)


# ---------------------------------------------------------------------------
# Attendance / occupancy / reports
# ---------------------------------------------------------------------------

@api_router.get("/attendance/inside")
async def currently_inside(_: dict = Depends(get_current_user)):
    cursor = db.attendance.find({"check_out_time": None}, {"_id": 0}).sort("check_in_time", -1)
    records = [d async for d in cursor]
    # Enrich with labour photo
    out = []
    for r in records:
        lab = await db.labours.find_one({"labour_id": r["labour_id"]}, {"_id": 0, "photo_base64": 1})
        r["photo_base64"] = lab["photo_base64"] if lab else None
        out.append(r)
    return out


@api_router.get("/attendance")
async def attendance_report(
    start: Optional[str] = None,
    end: Optional[str] = None,
    labour_id: Optional[str] = None,
    labour_name: Optional[str] = None,
    contractor_id: Optional[str] = None,
    category: Optional[str] = None,
    limit: int = Query(1000, le=5000),
    _: dict = Depends(get_current_user),
):
    flt: dict = {}
    s_dt = _parse_dt(start)
    e_dt = _parse_dt(end)
    if s_dt or e_dt:
        flt["check_in_time"] = {}
        if s_dt:
            flt["check_in_time"]["$gte"] = s_dt
        if e_dt:
            flt["check_in_time"]["$lte"] = e_dt
    if labour_id:
        flt["labour_id"] = {"$regex": re.escape(labour_id), "$options": "i"}
    if labour_name:
        flt["labour_name"] = {"$regex": re.escape(labour_name), "$options": "i"}
    if category:
        flt["category"] = category
    cursor = db.attendance.find(flt, {"_id": 0}).sort("check_in_time", -1).limit(limit)
    rows = [d async for d in cursor]
    if contractor_id:
        c = await db.contractors.find_one({"id": contractor_id}, {"_id": 0, "contractor_name": 1})
        if c:
            rows = [r for r in rows if r.get("contractor_name") == c["contractor_name"]]
    return rows


# ---------------------------------------------------------------------------
# Report exports (Excel + PDF)
# ---------------------------------------------------------------------------

def _parse_dt(v: Optional[str]) -> Optional[datetime]:
    if not v:
        return None
    dt = datetime.fromisoformat(v)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


def _fmt_dt(v) -> str:
    if not v:
        return ""
    if isinstance(v, str):
        try:
            v = datetime.fromisoformat(v.replace("Z", "+00:00"))
        except Exception:
            return v
    if v.tzinfo is None:
        v = v.replace(tzinfo=timezone.utc)
    return v.strftime("%Y-%m-%d %H:%M")


def _build_xlsx(headers: List[str], rows: List[List], title: str) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    # Header row
    ws.append(headers)
    head_fill = PatternFill(start_color="09090B", end_color="09090B", fill_type="solid")
    head_font = Font(bold=True, color="FFCC00")
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    # Auto width
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_pdf(headers: List[str], rows: List[List], title: str, subtitle: str = "") -> bytes:
    from io import BytesIO
    from reportlab.lib import colors as rl
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=title,
                            leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"<b>DBS FACTORY — {title}</b>", styles["Title"])]
    if subtitle:
        story.append(Paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 10))

    table_data = [headers] + [[str(c) if c is not None else "" for c in r] for r in rows]
    tbl = Table(table_data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rl.HexColor("#09090B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), rl.HexColor("#FFCC00")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl.HexColor("#FAFAFA"), rl.white]),
        ("GRID", (0, 0), (-1, -1), 0.4, rl.HexColor("#A1A1AA")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


def _export_response(data: bytes, filename: str, fmt: str) -> Response:
    media = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if fmt == "xlsx"
        else "application/pdf"
    )
    return Response(
        content=data,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/visitors/export")
async def export_visitors(
    format: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    start: Optional[str] = None,
    end: Optional[str] = None,
    _: dict = Depends(require_admin),
):
    flt: dict = {}
    s_dt = _parse_dt(start)
    e_dt = _parse_dt(end)
    if s_dt or e_dt:
        flt["entry_datetime"] = {}
        if s_dt:
            flt["entry_datetime"]["$gte"] = s_dt
        if e_dt:
            flt["entry_datetime"]["$lte"] = e_dt
    cursor = db.visitors.find(flt, {"_id": 0}).sort("entry_datetime", -1)
    headers = ["Visitor Name", "Mobile", "Purpose", "Gate", "Entry Date & Time"]
    rows = []
    async for d in cursor:
        rows.append([
            d.get("visitor_name", ""),
            d.get("mobile_number") or "",
            d.get("purpose", ""),
            d.get("gate_name") or "",
            _fmt_dt(d.get("entry_datetime")),
        ])
    sub = f"{start or 'all'} → {end or 'now'}  •  {len(rows)} records"
    if format == "xlsx":
        data = _build_xlsx(headers, rows, "Visitors")
    else:
        data = _build_pdf(headers, rows, "Visitor Report", sub)
    ts = utcnow().strftime("%Y%m%d_%H%M")
    return _export_response(data, f"visitors_{ts}.{format}", format)


@api_router.get("/attendance/export")
async def export_attendance(
    format: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    start: Optional[str] = None,
    end: Optional[str] = None,
    category: Optional[str] = None,
    _: dict = Depends(require_admin),
):
    flt: dict = {}
    s_dt = _parse_dt(start)
    e_dt = _parse_dt(end)
    if s_dt or e_dt:
        flt["check_in_time"] = {}
        if s_dt:
            flt["check_in_time"]["$gte"] = s_dt
        if e_dt:
            flt["check_in_time"]["$lte"] = e_dt
    if category:
        flt["category"] = category
    cursor = db.attendance.find(flt, {"_id": 0}).sort("check_in_time", -1)
    headers = ["Labour ID", "Name", "Contractor", "Category", "Check In", "Check Out", "Total Hours"]
    rows = []
    async for d in cursor:
        rows.append([
            d.get("labour_id", ""),
            d.get("labour_name", ""),
            d.get("contractor_name") or "",
            d.get("category", ""),
            _fmt_dt(d.get("check_in_time")),
            _fmt_dt(d.get("check_out_time")) if d.get("check_out_time") else "INSIDE",
            d.get("total_hours") if d.get("total_hours") is not None else "",
        ])
    sub = f"{start or 'all'} → {end or 'now'}  •  {len(rows)} records"
    if format == "xlsx":
        data = _build_xlsx(headers, rows, "Labour Attendance")
    else:
        data = _build_pdf(headers, rows, "Labour Attendance Report", sub)
    ts = utcnow().strftime("%Y%m%d_%H%M")
    return _export_response(data, f"labour_attendance_{ts}.{format}", format)


# ---------------------------------------------------------------------------
# Dashboard stats
# ---------------------------------------------------------------------------

def _start_of_today_utc() -> datetime:
    now = utcnow()
    return now.replace(hour=0, minute=0, second=0, microsecond=0)


@api_router.get("/dashboard/stats")
async def dashboard_stats(_: dict = Depends(require_admin)):
    sod = _start_of_today_utc()
    visitors_today = await db.visitors.count_documents({"entry_datetime": {"$gte": sod}})
    inside = await db.attendance.count_documents({"check_out_time": None})
    checked_out_today = await db.attendance.count_documents({
        "check_out_time": {"$gte": sod},
    })
    total_labour = await db.labours.count_documents({})
    contractor_count = await db.contractors.count_documents({})

    # per-contractor labour counts
    pipeline = [
        {"$group": {"_id": "$contractor_id", "count": {"$sum": 1}}},
    ]
    grouped = {}
    async for row in db.labours.aggregate(pipeline):
        grouped[row["_id"]] = row["count"]
    contractors = []
    async for c in db.contractors.find({}, {"_id": 0}):
        c["labour_count"] = grouped.get(c["id"], 0)
        # daily attendance count
        c["attendance_today"] = await db.attendance.count_documents({
            "check_in_time": {"$gte": sod},
            "contractor_name": c["contractor_name"],
        })
        contractors.append(c)

    return {
        "visitors_today": visitors_today,
        "labour_inside": inside,
        "checked_out_today": checked_out_today,
        "total_labour": total_labour,
        "contractor_count": contractor_count,
        "contractors": contractors,
    }


# ---------------------------------------------------------------------------
# Seed default admin on startup (idempotent)
# ---------------------------------------------------------------------------

@app.on_event("startup")
async def on_startup():
    await ensure_indexes()
    if not await db.users.find_one({"username": "admin"}):
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "name": "Administrator",
            "username": "admin",
            "password_hash": hash_password("admin123"),
            "role": "admin",
            "created_at": utcnow(),
        })
        logger.info("Seeded default admin user (admin / admin123)")
    if not await db.users.find_one({"username": "guard"}):
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "name": "Gate Guard",
            "username": "guard",
            "password_hash": hash_password("guard123"),
            "role": "guard",
            "created_at": utcnow(),
        })
        logger.info("Seeded default guard user (guard / guard123)")


@app.on_event("shutdown")
async def on_shutdown():
    client.close()


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)
